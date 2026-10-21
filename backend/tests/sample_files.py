"""Generators for real sample documents.

Fixtures are *generated* rather than committed as binaries. Two reasons: a binary
fixture is unreviewable in a pull request, and generating the file proves the app
handles output from the real libraries rather than from a hand-crafted byte string
that happens to parse.

Each generator writes a genuine file of its type, so the parsers under test do
real work.
"""

from __future__ import annotations

from pathlib import Path

HANDBOOK_PAGE_ONE = (
    "Employee Handbook\n\n"
    "Section 1: Annual leave\n"
    "Every full-time employee accrues twenty days of paid annual leave per calendar "
    "year. Leave requests must be submitted through the HR portal at least fourteen "
    "days before the planned absence. Unused leave may be carried over, up to a "
    "maximum of five days, into the following year."
)

HANDBOOK_PAGE_TWO = (
    "Section 2: Remote work\n"
    "Employees may work remotely up to three days per week with manager approval. "
    "Remote workers must remain reachable during core hours and attend all scheduled "
    "team meetings. Equipment provided by the company remains company property."
)

REVENUE_REPORT_TEXT = (
    "Q3 Revenue Report\n\n"
    "Total revenue for the third quarter reached 4,820,000 dollars, an increase of "
    "twelve percent over the previous quarter. The enterprise segment contributed "
    "3,100,000 dollars, driven by fourteen new annual contracts. The small business "
    "segment contributed 1,720,000 dollars. Gross margin held at sixty-one percent."
)

LAPTOP_SECURITY_TEXT = (
    "Device Security Standard\n\n"
    "All laptops issued to staff must have full-disk encryption enabled and a screen "
    "lock that engages after five minutes of inactivity. Encryption keys are escrowed "
    "with the platform team and rotated annually."
)


def write_text_file(
    directory: Path, name: str = "notes.txt", text: str = HANDBOOK_PAGE_ONE
) -> Path:
    path = directory / name
    path.write_text(text, encoding="utf-8")
    return path


def write_markdown_file(directory: Path, name: str = "guide.md") -> Path:
    path = directory / name
    path.write_text(
        "# Operations Guide\n\n"
        "## Incident response\n"
        "Severity one incidents must be acknowledged within fifteen minutes. The on-call "
        "engineer opens a bridge call and notifies the incident commander.\n\n"
        "## Change management\n"
        "Production changes require a peer review and a documented rollback plan.\n",
        encoding="utf-8",
    )
    return path


def write_csv_file(directory: Path, name: str = "sales.csv") -> Path:
    path = directory / name
    path.write_text(
        "region,product,units,revenue\n"
        "North,Widget,1200,48000\n"
        "South,Gadget,850,42500\n"
        "East,Widget,610,24400\n"
        "West,Doohickey,45,6750\n",
        encoding="utf-8",
    )
    return path


def write_xlsx_file(directory: Path, name: str = "budget.xlsx") -> Path:
    import pandas as pd

    path = directory / name
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame(
            {
                "department": ["Engineering", "Marketing", "Support"],
                "headcount": [24, 9, 12],
                "budget_usd": [1_800_000, 450_000, 320_000],
            }
        ).to_excel(writer, sheet_name="Annual Budget", index=False)
        pd.DataFrame(
            {"quarter": ["Q1", "Q2", "Q3"], "spend_usd": [600_000, 640_000, 705_000]}
        ).to_excel(writer, sheet_name="Spend", index=False)
    return path


def write_docx_file(directory: Path, name: str = "policy.docx") -> Path:
    from docx import Document

    path = directory / name
    document = Document()
    document.add_heading("Security Policy", level=1)
    document.add_paragraph(
        "Passwords must contain at least fourteen characters and are rotated every "
        "ninety days. Multi-factor authentication is mandatory for all staff accounts."
    )
    document.add_heading("Device encryption", level=2)
    document.add_paragraph(
        "All laptops must have full-disk encryption enabled before they are issued."
    )

    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Control"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Access review"
    table.cell(1, 1).text = "IT Security"
    table.cell(2, 0).text = "Backup testing"
    table.cell(2, 1).text = "Platform Team"

    document.save(str(path))
    return path


def write_pdf_file(
    directory: Path,
    name: str = "handbook.pdf",
    *,
    pages: int | None = None,
    texts: list[str] | None = None,
) -> Path:
    """Build a multi-page PDF with a real text layer.

    ``texts`` supplies one body per page; ``pages`` repeats the default set. Passing
    distinct text matters for retrieval tests - two documents with the same words
    score identically, so a ranking assertion on them would be decided by tie-break
    order rather than by similarity.
    """
    import pymupdf

    path = directory / name
    text_by_page = texts if texts else [HANDBOOK_PAGE_ONE, HANDBOOK_PAGE_TWO, REVENUE_REPORT_TEXT]
    page_count = pages if pages is not None else len(text_by_page)

    document = pymupdf.open()
    for index in range(page_count):
        page = document.new_page()
        body = text_by_page[index % len(text_by_page)]
        result = page.insert_textbox(
            pymupdf.Rect(50, 50, 545, 780),
            body,
            fontname="helv",
            fontsize=11,
        )
        if result < 0:
            raise RuntimeError(f"Sample PDF text did not fit on page {index + 1}")

    document.save(str(path))
    document.close()
    return path


def write_scanned_style_pdf(directory: Path, name: str = "scanned.pdf") -> Path:
    """A valid PDF with no text layer - stands in for a scanned image."""
    import pymupdf

    path = directory / name
    document = pymupdf.open()
    page = document.new_page()
    page.draw_rect(pymupdf.Rect(100, 100, 400, 300), color=(0, 0, 0), fill=(0.9, 0.9, 0.9))
    document.save(str(path))
    document.close()
    return path


def write_fake_pdf(directory: Path, name: str = "notreally.pdf") -> Path:
    """A file with a .pdf extension that is not a PDF (for validation tests)."""
    path = directory / name
    path.write_bytes(b"This is plain text wearing a PDF costume.\n" * 10)
    return path


def write_fake_docx(directory: Path, name: str = "notreally.docx") -> Path:
    """A genuine ZIP that is not a Word document (for magic-byte tests)."""
    import zipfile

    path = directory / name
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "just a zip, not a docx")
    return path
